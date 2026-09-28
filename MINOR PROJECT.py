# Generated from: MINOR PROJECT.ipynb
# Converted at: 2026-09-28T18:04:29.557Z
# Next step (optional): refactor into modules & generate tests with RunCell
# Quick start: pip install runcell

import openpyxl

# Dictionary mapping qualitative dropdown options to a 1-5 scale
mapping_scale = {
    # 1. Day's Feeling Options
    "excellent": 5,
    "good": 4,
    "neutral": 3,
    "low": 2,
    "stressed": 1,

    # 2. Satisfaction Level Options
    "very satisfied": 5,
    "satisfied": 4,
    "unsatisfied": 2,
    "very unsatisfied": 1,

    # 3. Energy Level Options
    "high": 5,
    "medium": 3,

    # 4. General Fallbacks
    "very poor": 1, "poor": 1, "bad": 1, "dissatisfied": 1,
    "below average": 2, "average": 3, "medium-high": 4
}


def parse_scale(val):
    """
    Converts qualitative dropdown text into a numeric 1-5 scale value.
    """
    if isinstance(val, (int, float)):
        return float(val)
    if isinstance(val, str):
        val_clean = val.strip().lower()
        return float(mapping_scale.get(val_clean, 3.0))
    return 3.0


def parse_num(val):
    """
    Safely converts cell values (including 'NA' or empty cells) to numeric floats.
    """
    if val is None or val == "" or str(val).strip().upper() == "NA":
        return 0.0
    try:
        return float(val)
    except (ValueError, TypeError):
        return 0.0


def scale_value(val, min_val, max_val):
    """
    Scales a value to a 0-1 range for the Personal Activity Index (PAI).
    """
    if max_val == min_val:
        return 0.0
    scaled = (val - min_val) / (max_val - min_val)
    return max(0.0, min(1.0, scaled))


def calculate_indices(file_path, sheet_name="Daily Log"):
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb[sheet_name]

        records = []

        # Iterate through data rows starting from row 6 (skipping headers/sample)
        for row in range(6, ws.max_row + 1):
            date_val = ws.cell(row=row, column=1).value
            
            # Skip empty date rows
            if date_val is None:
                continue

            # Read durations cleanly
            sleep = parse_num(ws.cell(row=row, column=2).value)
            fitness = parse_num(ws.cell(row=row, column=3).value)
            study = parse_num(ws.cell(row=row, column=4).value)
            coding = parse_num(ws.cell(row=row, column=5).value)
            class_min = parse_num(ws.cell(row=row, column=6).value)
            other_act = parse_num(ws.cell(row=row, column=8).value)
            total_tracked = parse_num(ws.cell(row=row, column=9).value)
            free_time = parse_num(ws.cell(row=row, column=10).value)

            # Read qualitative scales
            feeling = parse_scale(ws.cell(row=row, column=11).value)
            satisfaction = parse_scale(ws.cell(row=row, column=12).value)
            energy = parse_scale(ws.cell(row=row, column=13).value)

            records.append({
                'date': date_val,
                'sleep': sleep,
                'fitness': fitness,
                'study': study,
                'coding': coding,
                'class': class_min,
                'other': other_act,
                'total_tracked': total_tracked,
                'free_time': free_time,
                'feeling': feeling,
                'satisfaction': satisfaction,
                'energy': energy
            })

        valid_days = len(records)
        if valid_days == 0:
            print("No valid records found in the Excel file.")
            return

        # -------------------------------------------------------------
        # INDEX CALCULATIONS (Page 15)
        # -------------------------------------------------------------
        
        # 1. Tech Productivity (TPI)
        tpi = sum(r['coding'] for r in records) / valid_days

        # 2. Academic Activity (AAI)
        aai = sum(r['study'] + r['class'] for r in records) / valid_days

        # 3. Physical Activity (PhAI)
        phai = sum(r['fitness'] for r in records) / valid_days

        # 4. Sleep & Recovery Index (SRI)
        sri = sum(r['sleep'] for r in records) / valid_days

        # 5. Activity Balance Index (ABI)
        abi = sum(r['free_time'] for r in records) / valid_days

        # 6. Time Utilization Index (TUI)
        tui = sum(r['total_tracked'] for r in records) / valid_days

        # 7. Experience Index (EI)
        ei = sum(r['feeling'] + r['satisfaction'] + r['energy'] for r in records) / (3 * valid_days)

        # 8. Data Continuity Index (DCI) - Expected duration: 40 Days
        dci = (valid_days / 40) * 100

        # 9. Personal Activity Index (PAI)
        tpi_s = scale_value(tpi, 0, 180)
        aai_s = scale_value(aai, 0, 360)
        phai_s = scale_value(phai, 0, 60)
        sri_s = scale_value(sri, 0, 480)
        tui_s = scale_value(tui, 0, 1440)
        ei_s = scale_value(ei, 1, 5)

        pai = (0.15 * tpi_s + 0.20 * aai_s + 0.15 * phai_s +
               0.20 * sri_s + 0.15 * tui_s + 0.10 * ei_s + 0.05 * (dci / 100))

        # Console Display
        print("==========================================")
        print("          MY DATA, MY STORY RESULTS       ")
        print("==========================================")
        print(f"Valid Recorded Days   : {valid_days}")
        print(f"Tech Productivity (TPI): {tpi:.2f} min/day")
        print(f"Academic Activity (AAI): {aai:.2f} min/day")
        print(f"Physical Activity (PhAI): {phai:.2f} min/day")
        print(f"Sleep & Recovery (SRI) : {sri:.2f} min/day")
        print(f"Activity Balance (ABI) : {abi:.2f} min/day")
        print(f"Time Utilization (TUI) : {tui:.2f} min/day")
        print(f"Experience Index (EI)  : {ei:.2f} / 5")
        print(f"Data Continuity (DCI)  : {dci:.2f}%")
        print(f"Personal Activity Index: {pai:.2f}")
        print("==========================================")

        # Save results to output text file
        with open("Project_Summary_Report.txt", "w") as out_file:
            out_file.write("MY DATA, MY STORY - SUMMARY REPORT\n")
            out_file.write(f"Total Valid Days: {valid_days}\n")
            out_file.write(f"TPI: {tpi:.2f} min/day\n")
            out_file.write(f"AAI: {aai:.2f} min/day\n")
            out_file.write(f"PhAI: {phai:.2f} min/day\n")
            out_file.write(f"SRI: {sri:.2f} min/day\n")
            out_file.write(f"ABI: {abi:.2f} min/day\n")
            out_file.write(f"TUI: {tui:.2f} min/day\n")
            out_file.write(f"EI: {ei:.2f} / 5\n")
            out_file.write(f"DCI: {dci:.2f}%\n")
            out_file.write(f"PAI Score: {pai:.2f}\n")

        print("Summary report successfully saved to 'Project_Summary_Report.txt'.")

    except FileNotFoundError:
        print(f"Error: Could not find '{file_path}'. Ensure the file name is correct.")
    except Exception as e:
        print(f"An unexpected error occurred: {e}")


if __name__ == "__main__":
    calculate_indices(r"C:\Users\tania\OneDrive\Documents\CAP776\12616112.xlsx", sheet_name="Daily Log")